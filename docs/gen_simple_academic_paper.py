from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\nurym\Documents\AQA mid term\QA_drf-api")
DOCS = ROOT / "docs"
OUTDIR = ROOT / "output" / "doc"
OUTDIR.mkdir(parents=True, exist_ok=True)
OUTFILE = OUTDIR / "QA_Paper_Draft.docx"

TITLE = "A Literature-Grounded Evaluation of a Risk-Informed Multi-Level QA Strategy for an API-Centric E-Commerce Backend"

SECTION_FILES = [
    DOCS / "12_introduction_draft.md",
    DOCS / "13_lit_review_draft.md",
    DOCS / "14_method_draft.md",
    DOCS / "15_results_draft.md",
    DOCS / "16_discussion_draft.md",
]
REFERENCES_FILE = DOCS / "18_references_draft.md"

GRAPH_FILES = [
    (DOCS / "graph_coverage.png", "Figure 1. Coverage by high-risk module."),
    (DOCS / "graph_defects.png", "Figure 2. Defect distribution across modules."),
    (DOCS / "graph_runtime.png", "Figure 3. Runtime comparison by test layer."),
]


def set_cell_shading(cell, fill: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_font(run, size: int, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def apply_paragraph_style(paragraph, size: int = 12, bold: bool = False, italic: bool = False) -> None:
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, italic=italic)


def add_body_paragraph(doc: Document, text: str, indent: bool = True, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5) if indent else Inches(0)
    p.add_run(text)
    apply_paragraph_style(p, size=12, italic=italic)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, bold=True)
    elif level == 2:
        set_run_font(run, size=14, bold=True)
    else:
        set_run_font(run, size=12, bold=True)


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("`", "")
    text = text.replace("\u2014", "-")
    text = text.replace("\u2013", "-")
    text = re.sub(r"\s+Draft$", "", text)
    return text.strip()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    block = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        block.append(lines[i].rstrip())
        i += 1

    rows = []
    for idx, line in enumerate(block):
        if idx == 1 and set(line.replace("|", "").replace("-", "").replace(":", "").strip()) == set():
            continue
        parts = [clean_inline(cell.strip()) for cell in line.strip().strip("|").split("|")]
        rows.append(parts)
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[c_idx] if c_idx < len(row) else ""
            cell.text = text
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, size=11, bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.5
    p.add_run(clean_inline(text))
    for run in p.runs:
        set_run_font(run, size=12)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    p.add_run(text)
    apply_paragraph_style(p, size=11)


def add_figure_block(doc: Document) -> None:
    for path, caption in GRAPH_FILES:
        if path.exists():
            doc.add_picture(str(path), width=Inches(6.0))
            add_caption(doc, caption)


def add_references(doc: Document) -> None:
    add_heading(doc, "References", level=1)
    for line in REFERENCES_FILE.read_text(encoding="utf-8").splitlines():
        text = line.strip()
        if not text or text.startswith("#"):
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(3)
        p.add_run(text)
        apply_paragraph_style(p, size=12)


def process_markdown_file(doc: Document, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    skip_visual_block = False

    while i < len(lines):
        raw = lines[i].rstrip()
        line = raw.strip()

        if not line or line == "---":
            i += 1
            continue

        if path.name == "15_results_draft.md" and line.startswith("For the article version, these results should be accompanied by"):
            add_figure_block(doc)
            skip_visual_block = True
            i += 1
            continue

        if skip_visual_block:
            if line.startswith("## 3."):
                skip_visual_block = False
            else:
                i += 1
                continue

        if path.name == "15_results_draft.md" and "For the article version, this should be presented as a compact screenshot-based figure" in line:
            i += 1
            continue

        if line.startswith("# "):
            add_heading(doc, clean_inline(line[2:]), level=1)
            i += 1
            continue

        if line.startswith("## "):
            add_heading(doc, clean_inline(line[3:]), level=2)
            i += 1
            continue

        if line.startswith("### "):
            add_heading(doc, clean_inline(line[4:]), level=3)
            i += 1
            continue

        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        if line.startswith("- "):
            add_bullet(doc, line[2:])
            i += 1
            continue

        add_body_paragraph(doc, clean_inline(line))
        i += 1


def build_doc() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    title_run = title.add_run(TITLE)
    set_run_font(title_run, size=16, bold=True)

    abs_heading = doc.add_paragraph()
    abs_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    abs_run = abs_heading.add_run("Abstract")
    set_run_font(abs_run, size=16, bold=True)

    add_body_paragraph(
        doc,
        "Abstract will be completed in the final version of the paper after the full article text and final results framing are stabilized.",
        indent=False,
        italic=True,
    )

    for path in SECTION_FILES:
        process_markdown_file(doc, path)

    add_references(doc)

    doc.save(OUTFILE)


if __name__ == "__main__":
    build_doc()
    print(OUTFILE)
